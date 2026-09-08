"""
excel_exporter.py
000_Form.xlsx 템플릿의 오류 수식을 정정하고, 수집된 데이터를 기록하여
수식이 온전히 작동하는 최종 Excel 파일(.xlsx)을 생성하는 모듈
"""

import os
import re
import io
import zipfile
import shutil
from datetime import datetime
import openpyxl
from openpyxl.utils import get_column_letter

def fix_xlsx_relations_bytes(excel_bytes):
    """
    openpyxl의 버그로 인해 내부 XML relationship(.rels)의 경로가 절대경로(/xl/...)로
    작성되어 엑셀 복구 오류가 발생하는 문제를 메모리 상에서 상대경로로 수정
    """
    in_buf = io.BytesIO(excel_bytes)
    out_buf = io.BytesIO()
    with zipfile.ZipFile(in_buf, 'r') as yin:
        with zipfile.ZipFile(out_buf, 'w', zipfile.ZIP_DEFLATED) as yout:
            for item in yin.infolist():
                content = yin.read(item.filename)
                if item.filename.endswith('.rels'):
                    text = content.decode('utf-8', errors='ignore')
                    if item.filename == 'xl/_rels/workbook.xml.rels':
                        text = re.sub(r'Target="/xl/([^"]+)"', r'Target="\1"', text)
                    elif 'xl/worksheets/_rels/' in item.filename:
                        text = re.sub(r'Target="/xl/([^"]+)"', r'Target="../\1"', text)
                    elif 'xl/drawings/_rels/' in item.filename:
                        text = re.sub(r'Target="/xl/([^"]+)"', r'Target="../\1"', text)
                    content = text.encode('utf-8')
                yout.writestr(item, content)
    return out_buf.getvalue()

def patch_template_formulas(wb):
    """
    기존 000_Form.xlsx 템플릿에 존재하던 수식 오류들을 교정
    """
    # 1. RIM계산 시트 B3 셀의 요구수익률 참조 단절 오류 수정
    if "RIM계산" in wb.sheetnames:
        ws_rim = wb["RIM계산"]
        # 기존: ='S-RIM'!D8 (하드코딩된 별도요구수익률만 참조)
        # 수정: ='S-RIM'!C11 (S-RIM 시트의 최종 적용 할인율 참조)
        ws_rim["B3"] = "='S-RIM'!C11"

    # 2. Q 시트 O20 셀의 0으로 나누기 오류 수정
    if "Q" in wb.sheetnames:
        ws_q = wb["Q"]
        # 기존: =IF(O15=0, O19/O15, O17/O19)
        # 수정: =IF(OR(O15=0, O19=0), 0, O17/O19)
        ws_q["O20"] = "=IF(OR(O15=0, O19=0), 0, O17/O19)"

    # 3. S-RIM 시트 G18 셀의 컨센서스 결측 검사 수정
    if "S-RIM" in wb.sheetnames:
        ws_srim = wb["S-RIM"]
        # 기존: =IF(E18="컨센 없음", IF(D18="컨센 없음", IF(C18="컨센 없음", ...)))
        # 빈 셀(/100 시 0)도 컨센서스 없음으로 인식하도록 안전하게 수정
        orig_formula = str(ws_srim["G18"].value or "")
        if "컨센 없음" in orig_formula:
            # 안전한 수식으로 교체
            ws_srim["G18"] = '=IF(OR(E18="컨센 없음", E18=0, ISBLANK(Data!I62)), IF(OR(D18="컨센 없음", D18=0, ISBLANK(Data!H62)), IF(OR(C18="컨센 없음", C18=0, ISBLANK(Data!G62)), IF(H16="분기자료부족", J16, IF(I15="상승추세", H16, IF(I16="하락추세", H16, F18))), C18), D18), E18)'

def write_val(ws, row, col, val):
    """셀에 값이 수식이 아닌 경우에만 안전하게 기록"""
    cell = ws.cell(row=row, column=col)
    if cell.value and str(cell.value).startswith("="):
        return
    cell.value = val

def generate_srim_excel(scraped_data, form_path="000_Form.xlsx"):
    """
    scraped_data(fnGuide 스크래핑 결과)를 바탕으로
    수식과 서식이 온전히 보존된 .xlsx 바이트 데이터를 생성
    """
    if not os.path.exists(form_path):
        raise FileNotFoundError(f"템플릿 파일이 존재하지 않습니다: {form_path}")

    wb = openpyxl.load_workbook(form_path, data_only=False)
    patch_template_formulas(wb)

    ws_data = wb["Data"]
    info = scraped_data.get('info', {})
    ticker = scraped_data.get('ticker', '')
    company_name = scraped_data.get('company_name', '')
    inquiry_date = scraped_data.get('inquiry_date', datetime.now().strftime("%Y/%m/%d"))

    # 1. 기본 정보 기록
    write_val(ws_data, 3, 2, company_name)
    write_val(ws_data, 10, 2, inquiry_date)
    write_val(ws_data, 24, 2, inquiry_date)

    # 밸류에이션 비율 (Rows 4-8)
    write_val(ws_data, 4, 2, info.get('per'))
    write_val(ws_data, 5, 2, info.get('fwd_per'))
    write_val(ws_data, 6, 2, info.get('industry_per'))
    write_val(ws_data, 7, 2, info.get('pbr'))
    write_val(ws_data, 8, 2, info.get('dividend_yield'))

    # 시세 현황 (Rows 11-19)
    write_val(ws_data, 11, 2, info.get('current_price'))
    write_val(ws_data, 12, 2, info.get('high_52'))
    write_val(ws_data, 13, 2, info.get('low_52'))
    write_val(ws_data, 14, 2, info.get('market_cap_total'))
    write_val(ws_data, 15, 2, info.get('market_cap_common'))
    write_val(ws_data, 16, 2, info.get('shares_common'))
    write_val(ws_data, 17, 2, info.get('shares_pref'))
    write_val(ws_data, 19, 2, info.get('beta'))

    # 자기주식 (Row 22)
    write_val(ws_data, 22, 2, info.get('shares_treasury'))

    # 투자의견 요약 (Rows 25-29)
    cons_info = scraped_data.get('consensus_info', {})
    write_val(ws_data, 25, 2, cons_info.get('opinion_score'))
    write_val(ws_data, 26, 2, cons_info.get('target_price_avg'))
    write_val(ws_data, 27, 2, cons_info.get('eps'))
    write_val(ws_data, 28, 2, cons_info.get('per'))
    write_val(ws_data, 29, 2, cons_info.get('analyst_count'))

    # 업종비교 (Rows 33-38)
    ind = scraped_data.get('industry_comp', {})
    if ind:
        write_val(ws_data, 33, 2, ind.get('company_name'))
        write_val(ws_data, 33, 3, ind.get('industry_name'))
        write_val(ws_data, 33, 4, ind.get('market_name'))
        metrics = ind.get('metrics', {})
        term_rows = [(34, "PER"), (35, "EV/EBITDA"), (36, "ROE"), (37, "배당수익률"), (38, "베타(1년)")]
        for r_idx, t_name in term_rows:
            m_data = metrics.get(t_name) or metrics.get(t_name.replace(" ", ""))
            if m_data:
                write_val(ws_data, r_idx, 2, m_data.get('company'))
                write_val(ws_data, r_idx, 3, m_data.get('industry'))
                write_val(ws_data, r_idx, 4, m_data.get('market'))

    # 2. Financial Highlights - 연간 (Rows 44-69)
    df_ann = scraped_data.get('highlights_annual')
    if df_ann is not None and not df_ann.empty:
        # 헤더 (B44 ~ I44)
        for c_idx, col_name in enumerate(df_ann.columns[:8], start=2):
            write_val(ws_data, 44, c_idx, str(col_name))
            
        # 값 매핑
        for r_idx in range(45, 70):
            label = ws_data.cell(row=r_idx, column=1).value
            if not label:
                continue
            clean_l = str(label).replace("\xa0", " ").replace(" ", "").strip()
            # df_ann에서 가장 잘 매칭되는 행 탐색
            matched_row = None
            for df_idx in df_ann.index:
                clean_df_idx = str(df_idx).replace("\xa0", " ").replace(" ", "").strip()
                if clean_l == clean_df_idx or clean_l in clean_df_idx or clean_df_idx in clean_l:
                    matched_row = df_ann.loc[df_idx]
                    break
            if matched_row is not None:
                for c_idx, col_name in enumerate(df_ann.columns[:8], start=2):
                    val = matched_row.get(col_name)
                    write_val(ws_data, r_idx, c_idx, val)

    # 3. Financial Highlights - 분기 (Rows 73-98)
    df_qtr = scraped_data.get('highlights_quarter')
    if df_qtr is not None and not df_qtr.empty:
        for c_idx, col_name in enumerate(df_qtr.columns[:8], start=2):
            write_val(ws_data, 73, c_idx, str(col_name))
            
        for r_idx in range(74, 99):
            label = ws_data.cell(row=r_idx, column=1).value
            if not label:
                continue
            clean_l = str(label).replace("\xa0", " ").replace(" ", "").strip()
            matched_row = None
            for df_idx in df_qtr.index:
                clean_df_idx = str(df_idx).replace("\xa0", " ").replace(" ", "").strip()
                if clean_l == clean_df_idx or clean_l in clean_df_idx or clean_df_idx in clean_l:
                    matched_row = df_qtr.loc[df_idx]
                    break
            if matched_row is not None:
                for c_idx, col_name in enumerate(df_qtr.columns[:8], start=2):
                    val = matched_row.get(col_name)
                    write_val(ws_data, r_idx, c_idx, val)

    # 4. Consensus 상세 (Rows 193-205)
    reports = scraped_data.get('consensus_reports', [])
    if reports:
        # Row 195: 평균
        if cons_info.get('target_price_avg'):
            write_val(ws_data, 195, 3, cons_info.get('target_price_avg'))
            write_val(ws_data, 195, 6, cons_info.get('opinion_score'))
            
        # Row 196 ~ 205 (최대 10개 증권사)
        for idx, rep in enumerate(reports[:10]):
            r_idx = 196 + idx
            write_val(ws_data, r_idx, 1, rep.get('broker'))
            write_val(ws_data, r_idx, 2, rep.get('date'))
            write_val(ws_data, r_idx, 3, rep.get('target_price'))
            write_val(ws_data, r_idx, 4, rep.get('prev_price'))
            write_val(ws_data, r_idx, 5, rep.get('change_rate'))
            write_val(ws_data, r_idx, 6, rep.get('opinion'))

    # 메모리 버퍼에 저장
    temp_buf = io.BytesIO()
    if hasattr(wb, '_external_links'):
        wb._external_links = []
    wb.save(temp_buf)
    
    # openpyxl rels 버그 수정
    fixed_bytes = fix_xlsx_relations_bytes(temp_buf.getvalue())
    return fixed_bytes

if __name__ == "__main__":
    from fnguide_scraper import scrape_company_data
    print("Testing excel_exporter on 005930...")
    data = scrape_company_data("005930")
    excel_bytes = generate_srim_excel(data, form_path="000_Form.xlsx")
    test_out = "test_srim_out.xlsx"
    with open(test_out, "wb") as f:
        f.write(excel_bytes)
    print(f"Success! Generated {test_out} ({len(excel_bytes):,} bytes)")
